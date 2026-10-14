"""Data-linkage and pixel hazards; optional dependencies are skipped explicitly."""

import csv
import importlib.util
import json
from pathlib import Path
import tempfile
import unittest

from breast_unet.data_audit import (
    IMAGE_DIRECTORY, WORKBOOK_FILENAME, audit_dataset, inspect_png,
    match_rows, parse_filenames,
)

HAS_IMAGES = importlib.util.find_spec("numpy") is not None and importlib.util.find_spec("PIL") is not None
HAS_WORKBOOK = importlib.util.find_spec("openpyxl") is not None


def metadata(case_id=1, classification="benign", tumor="case001_tumor.png", other=None):
    return {
        "CaseID": case_id, "Image_filename": f"case{case_id:03d}.png",
        "Mask_tumor_filename": tumor, "Mask_other_filename": other,
        "Classification": classification, "Margin": "circumscribed",
    }


class MetadataMatchingTests(unittest.TestCase):
    def test_missing_tumor_is_not_normal(self):
        cases, _ = match_rows([metadata(tumor=None)])
        self.assertEqual(cases[0]["candidate_cohort"], "unresolved")
        self.assertIn("lesion_case_missing_tumor_annotation", cases[0]["issues"])
        self.assertEqual(cases[0]["classification"], "benign")

    def test_normal_is_based_on_classification_and_not_training_ready(self):
        cases, _ = match_rows([metadata(classification="normal", tumor=None)])
        self.assertEqual(cases[0]["candidate_cohort"], "normal_exploration_pending_review")
        self.assertEqual(cases[0]["issues"], [])
        self.assertFalse(cases[0]["training_ready"])
        self.assertEqual(cases[0]["visual_review_status"], "pending_visual_review")

    def test_normal_with_tumor_is_a_conflict(self):
        cases, _ = match_rows([metadata(classification="normal")])
        self.assertIn("normal_classification_with_tumor_annotation", cases[0]["issues"])

    def test_other_masks_never_replace_missing_tumor(self):
        cases, references = match_rows([metadata(tumor=None, other="case001_other1.png&case001_other2.png")])
        self.assertEqual(cases[0]["tumor_filenames"], [])
        self.assertEqual(len(cases[0]["other_filenames"]), 2)
        self.assertIn("lesion_case_missing_tumor_annotation", cases[0]["issues"])
        self.assertEqual(references["case001_other1.png"][0]["role"], "other")

    def test_multiple_tumors_remain_separate_and_require_review(self):
        cases, _ = match_rows([metadata(tumor="case001_tumor1.png&case001_tumor2.png")])
        self.assertEqual(len(cases[0]["tumor_filenames"]), 2)
        self.assertIn("multiple_tumor_masks_not_merged", cases[0]["review_flags"])

    def test_rejects_dangerous_or_duplicate_filenames(self):
        for value in ("../case001.png", "/case001.png", "x\\case001.png", "C:case001.png", "case001.png&", "case001.png&case001.png", "case001.png\n"):
            with self.subTest(value=value):
                if value.endswith("\n"):
                    # Surrounding whitespace is harmless and stripped.
                    self.assertEqual(parse_filenames(value), ["case001.png"])
                else:
                    with self.assertRaises(ValueError):
                        parse_filenames(value)

    def test_duplicate_case_and_file_mapping_are_blocked(self):
        cases, _ = match_rows([metadata(), metadata()])
        for case in cases:
            self.assertIn("duplicate_case_id", case["issues"])
            self.assertIn("duplicate_file_mapping", case["issues"])

    def test_swapped_mask_roles_and_wrong_case_are_flagged(self):
        cases, _ = match_rows([metadata(tumor="case001_other.png", other="case002_tumor.png")])
        self.assertIn("tumor_filename_case_or_role_mismatch", cases[0]["issues"])
        self.assertIn("other_filename_case_or_role_mismatch", cases[0]["issues"])


@unittest.skipUnless(HAS_IMAGES, "numpy and Pillow are optional audit dependencies")
class PixelInspectionTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.path = Path(self.tmp.name) / "mask.png"

    def tearDown(self):
        self.tmp.cleanup()

    def save(self, array):
        from PIL import Image
        Image.fromarray(array).save(self.path)

    def test_opaque_alpha_does_not_make_whole_mask_foreground(self):
        import numpy as np
        array = np.zeros((4, 4, 4), dtype=np.uint8)
        array[..., 3] = 255
        array[1, 1, :3] = 255
        self.save(array)
        result = inspect_png(self.path, "tumor")
        self.assertEqual(result["foreground_pixels"], 1)
        self.assertEqual(result["alpha_handling"], "verified_opaque_255_then_ignored")
        self.assertEqual(result["issues"], [])

    def test_nonopaque_alpha_is_rejected(self):
        import numpy as np
        array = np.full((4, 4, 4), 255, dtype=np.uint8)
        array[0, 0, 3] = 0
        self.save(array)
        result = inspect_png(self.path, "tumor")
        self.assertIn("nonopaque_alpha", result["issues"])
        self.assertNotIn("foreground_pixels", result)
        self.assertNotIn("foreground_pixels_256", result)
        self.assertEqual(result["alpha_zero_pixels"], 1)

    def test_color_content_is_not_silently_converted(self):
        import numpy as np
        array = np.zeros((4, 4, 3), dtype=np.uint8)
        array[0, 0, 0] = 255
        self.save(array)
        self.assertIn("unequal_color_channels", inspect_png(self.path, "image")["issues"])

    def test_empty_nonbinary_and_nearest_resize_disappearance(self):
        import numpy as np
        for array, expected in (
            (np.zeros((4, 4), dtype=np.uint8), "empty_mask"),
            (np.array([[0, 127], [255, 0]], dtype=np.uint8), "nonbinary_mask"),
        ):
            with self.subTest(expected=expected):
                self.save(array)
                self.assertIn(expected, inspect_png(self.path, "tumor")["issues"])
        tiny = np.zeros((512, 512), dtype=np.uint8)
        tiny[0, 0] = 255
        self.save(tiny)
        self.assertIn("mask_disappears_at_256", inspect_png(self.path, "tumor")["issues"])

    def test_grayscale_duplicates_detect_equal_pixels_across_modes(self):
        import numpy as np
        array = np.arange(16, dtype=np.uint8).reshape(4, 4)
        self.save(array)
        first = inspect_png(self.path, "image")
        self.save(np.repeat(array[..., None], 3, axis=2))
        second = inspect_png(self.path, "image")
        self.assertEqual(first["grayscale_sha256"], second["grayscale_sha256"])
        self.assertNotEqual(first["decoded_sha256"], second["decoded_sha256"])


@unittest.skipUnless(HAS_IMAGES and HAS_WORKBOOK, "numpy, Pillow and openpyxl are optional audit dependencies")
class DatasetAuditTests(unittest.TestCase):
    def test_empty_workbook_is_rejected_without_output(self):
        from openpyxl import Workbook

        with tempfile.TemporaryDirectory() as directory:
            project = Path(directory)
            raw = project / "data/raw/fixture"
            (raw / IMAGE_DIRECTORY).mkdir(parents=True)
            workbook = Workbook()
            workbook.active.append(list(metadata()))
            workbook.save(raw / WORKBOOK_FILENAME)
            out = project / "data/audit/empty"
            with self.assertRaisesRegex(ValueError, "no case records"):
                audit_dataset(project, raw, out)
            self.assertFalse(out.exists())

    def test_decoded_color_duplicates_and_mask_coverage(self):
        import numpy as np
        from PIL import Image
        from openpyxl import Workbook

        with tempfile.TemporaryDirectory() as directory:
            project = Path(directory)
            raw = project / "data/raw/fixture"
            images = raw / IMAGE_DIRECTORY
            images.mkdir(parents=True)
            workbook = Workbook()
            workbook.active.append(list(metadata()))
            array = np.zeros((4, 4, 4), dtype=np.uint8)
            array[..., 3] = 255
            array[..., 0] = 128  # Color prevents automatic gray extraction.
            mask = np.zeros((4, 4, 4), dtype=np.uint8)
            mask[1, 1, :] = 255  # Transparent background remains unresolved.
            for case_id in (1, 2):
                row = metadata(case_id=case_id, tumor=f"case{case_id:03d}_tumor.png")
                workbook.active.append(list(row.values()))
                # Different compression still produces equal decoded pixels.
                Image.fromarray(array).save(images / row["Image_filename"], compress_level=case_id)
                Image.fromarray(mask).save(images / row["Mask_tumor_filename"])
            workbook.save(raw / WORKBOOK_FILENAME)
            out = project / "data/audit/duplicates"
            summary = audit_dataset(project, raw, out)
            self.assertEqual(len(summary["duplicate_image_decoded_groups"]), 1)
            self.assertEqual(summary["review_flag_counts"]["duplicate_image_content"], 2)
            self.assertEqual(summary["mask_check_coverage"]["masks_total"], 2)
            self.assertEqual(summary["mask_check_coverage"]["masks_with_validated_binary_values"], 0)
            self.assertEqual(summary["mask_check_coverage"]["masks_resize_checked"], 0)
            self.assertEqual(summary["mask_check_coverage"]["masks_not_resize_checked"], 2)
            self.assertEqual(summary["masks_disappearing_at_256"], [])
            self.assertFalse(summary["automatic_checks_passed"])

    def test_dimension_missing_file_orphan_privacy_and_no_overwrite(self):
        import numpy as np
        from PIL import Image
        from openpyxl import Workbook

        with tempfile.TemporaryDirectory() as directory:
            project = Path(directory)
            raw = project / "data/raw/fixture"
            images = raw / IMAGE_DIRECTORY
            images.mkdir(parents=True)
            row = metadata(other="case001_other.png")
            row["Age"] = 99  # This unnecessary field must not reach the outputs.
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(list(row))
            sheet.append(list(row.values()))
            workbook.save(raw / WORKBOOK_FILENAME)
            Image.fromarray(np.arange(16, dtype=np.uint8).reshape(4, 4)).save(images / "case001.png")
            Image.fromarray(np.array([[255, 0], [0, 0]], dtype=np.uint8)).save(images / "case001_tumor.png")
            Image.fromarray(np.zeros((2, 2), dtype=np.uint8)).save(images / "orphan.png")
            out = project / "data/audit/test"
            summary = audit_dataset(project, raw, out)
            self.assertFalse(summary["automatic_checks_passed"])
            self.assertEqual(summary["case_issue_counts"]["missing_other_file"], 1)
            self.assertEqual(summary["case_issue_counts"]["tumor:image_mask_dimension_mismatch"], 1)
            self.assertEqual(len(summary["orphan_png_paths"]), 1)
            self.assertFalse(summary["training_ready"])
            with (out / "cases.csv").open(encoding="utf-8-sig", newline="") as handle:
                cases = list(csv.DictReader(handle))
            self.assertNotIn("Age", cases[0])
            self.assertEqual(cases[0]["visual_review_status"], "pending_visual_review")
            for path in out.iterdir():
                self.assertNotIn(str(project), path.read_text(encoding="utf-8-sig"))
            with self.assertRaises(FileExistsError):
                audit_dataset(project, raw, out)


if __name__ == "__main__":
    unittest.main()
