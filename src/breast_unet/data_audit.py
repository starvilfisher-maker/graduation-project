"""First-pass BrEaST audit; no split, augmentation, mask union, or training.

Metadata matching is standard-library-only. Image and Excel dependencies are
loaded only by the corresponding operations so base project tests stay offline.
Every case remains pending visual review even if all automatic checks pass.
"""

from collections import Counter, defaultdict
import csv
import hashlib
import json
from pathlib import Path
import platform
import re
from typing import Any


IMAGE_DIRECTORY = "BrEaST-Lesions_USG-images_and_masks"
WORKBOOK_FILENAME = "BrEaST-Lesions-USG-clinical-data-Dec-15-2023.xlsx"
REQUIRED_COLUMNS = (
    "CaseID", "Image_filename", "Mask_tumor_filename",
    "Mask_other_filename", "Classification", "Margin",
)
PATIENT_ID_BASIS = (
    "patient_id is derived from the workbook CaseID for record linkage, based "
    "on the published BrEaST description of one image per participant. It is "
    "not a separately verified clinical identity; file/pixel duplicate checks "
    "cannot establish whether two different images show the same participant."
)


def parse_filenames(value: Any) -> list[str]:
    """Parse explicit ampersand-separated PNG basenames; never infer a mask."""
    if value is None or value == "":
        return []
    if not isinstance(value, str):
        raise ValueError("filenames must be strings")
    if not value.strip():
        return []
    names = [part.strip() for part in value.split("&")]
    for name in names:
        if (not name or name in {".", ".."} or "/" in name or "\\" in name
                or ":" in name or any(ord(char) < 32 for char in name)
                or not name.lower().endswith(".png")):
            raise ValueError("unsafe or invalid PNG basename")
    if len(set(names)) != len(names):
        raise ValueError("duplicate filename within one field")
    return names


def normalize_case_id(value: Any) -> int:
    if isinstance(value, bool):
        raise ValueError("CaseID must be a positive integer")
    if isinstance(value, int):
        result = value
    elif isinstance(value, float) and value.is_integer():
        result = int(value)
    elif isinstance(value, str) and re.fullmatch(r"[0-9]+", value.strip()):
        result = int(value)
    else:
        raise ValueError("CaseID must be a positive integer")
    if result <= 0:
        raise ValueError("CaseID must be positive")
    return result


def match_rows(rows: list[dict[str, Any]]) -> tuple[list[dict], dict[str, list]]:
    """Keep tumor/other separate; missing tumor does not imply normality."""
    cases = []
    references: dict[str, list] = defaultdict(list)
    id_rows: dict[int, list[int]] = defaultdict(list)
    for offset, row in enumerate(rows):
        case = {
            "source_row": row.get("__source_row__", offset + 2),
            "case_id": "", "patient_id": "", "image_id": "",
            "classification": str(row.get("Classification") or "").strip().lower(),
            "margin": str(row.get("Margin") or "").strip(),
            "image_filenames": [], "tumor_filenames": [], "other_filenames": [],
            "candidate_cohort": "unresolved",
            "issues": [], "review_flags": [],
            "visual_review_status": "pending_visual_review", "training_ready": False,
        }
        try:
            case_id = normalize_case_id(row.get("CaseID"))
            case["case_id"] = case_id
            case["patient_id"] = case["image_id"] = f"case{case_id:03d}"
            id_rows[case_id].append(offset)
        except ValueError:
            case["issues"].append("invalid_case_id")
        for role, column in (
            ("image", "Image_filename"), ("tumor", "Mask_tumor_filename"),
            ("other", "Mask_other_filename"),
        ):
            try:
                filenames = parse_filenames(row.get(column))
            except ValueError:
                case["issues"].append(f"invalid_{role}_filename_field")
                filenames = []
            case[f"{role}_filenames"] = filenames
            for name in filenames:
                references[name].append({"row_index": offset, "role": role})
                if case["case_id"]:
                    prefix = f"case{case['case_id']:03d}"
                    suffix = r"\.png" if role == "image" else rf"_{role}(?:[1-9][0-9]*)?\.png"
                    if not re.fullmatch(re.escape(prefix) + suffix, name):
                        case["issues"].append(f"{role}_filename_case_or_role_mismatch")
        if len(case["image_filenames"]) != 1:
            case["issues"].append("expected_exactly_one_image")
        classification = case["classification"]
        tumor_count = len(case["tumor_filenames"])
        if classification in {"benign", "malignant"}:
            if tumor_count:
                case["candidate_cohort"] = "lesion_positive_pending_review"
            else:
                case["issues"].append("lesion_case_missing_tumor_annotation")
        elif classification == "normal":
            if tumor_count:
                case["issues"].append("normal_classification_with_tumor_annotation")
            else:
                case["candidate_cohort"] = "normal_exploration_pending_review"
        else:
            case["issues"].append("unknown_classification")
        if tumor_count > 1:
            case["review_flags"].append("multiple_tumor_masks_not_merged")
        cases.append(case)
    for indices in id_rows.values():
        if len(indices) > 1:
            for index in indices:
                cases[index]["issues"].append("duplicate_case_id")
    for occurrences in references.values():
        if len(occurrences) > 1:
            for item in occurrences:
                cases[item["row_index"]]["issues"].append("duplicate_file_mapping")
    return cases, dict(references)


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def inspect_png(path: Path, role: str) -> dict:
    """Inspect without altering originals; opaque alpha is explicitly ignored.

    Mask foreground is decoded from the verified grayscale channel only, never
    from alpha. Supported binary encodings are 0/1 and 0/255. Other encodings,
    nonopaque alpha, and unequal color channels are flagged, not guessed.
    """
    import numpy as np
    from PIL import Image

    result: dict[str, Any] = {"issues": [], "role": role}
    try:
        with Image.open(path) as image:
            image.load()
            result.update(width=image.width, height=image.height, mode=image.mode)
            array = np.asarray(image)
            result["dtype"] = str(array.dtype)
            digest = hashlib.sha256()
            digest.update(json.dumps([image.mode, list(array.shape), str(array.dtype)]).encode())
            digest.update(array.tobytes())
            result["decoded_sha256"] = digest.hexdigest()
            result["alpha_handling"] = "absent"
            if image.mode in {"RGB", "RGBA"}:
                rgb = array[..., :3]
                result["rgb_channels_equal"] = bool(np.array_equal(rgb[..., 0], rgb[..., 1]) and np.array_equal(rgb[..., 0], rgb[..., 2]))
                result["rgb_values_only_0_255"] = bool(np.all((rgb == 0) | (rgb == 255)))
            if image.mode in {"RGBA", "LA"}:
                alpha = array[..., -1]
                result["alpha_min"] = int(alpha.min())
                result["alpha_max"] = int(alpha.max())
                result["alpha_zero_pixels"] = int((alpha == 0).sum())
                result["alpha_partial_pixels"] = int(((alpha > 0) & (alpha < 255)).sum())
                result["transparent_color_all_zero"] = bool(np.all(array[..., :-1][alpha == 0] == 0))
                result["alpha_equals_first_channel"] = bool(np.array_equal(alpha, array[..., 0]))
                result["opaque_color_all_255"] = bool(np.all(array[..., :-1][alpha == 255] == 255))
                if not np.all(alpha == 255):
                    result["issues"].append("nonopaque_alpha")
                    result["alpha_handling"] = "rejected_nonopaque"
                    return result
                result["alpha_handling"] = "verified_opaque_255_then_ignored"
                array = array[..., :-1]
            if image.mode not in {"1", "L", "I", "I;16", "I;16B", "I;16L", "RGB", "RGBA", "LA"}:
                result["issues"].append("unsupported_image_mode")
                return result
            if array.ndim == 3:
                if array.shape[2] == 1:
                    array = array[..., 0]
                elif array.shape[2] == 3 and np.array_equal(array[..., 0], array[..., 1]) and np.array_equal(array[..., 0], array[..., 2]):
                    array = array[..., 0]
                else:
                    result["issues"].append("unequal_color_channels")
                    return result
            if array.ndim != 2:
                result["issues"].append("unexpected_pixel_shape")
                return result
            result["grayscale_min"] = int(array.min())
            result["grayscale_max"] = int(array.max())
            if result["grayscale_min"] < 0:
                result["issues"].append("negative_pixel_values")
            # Canonical grayscale content detects identical pixels even when
            # PNG mode, channel repetition, or compression differs.
            gray_hash = hashlib.sha256()
            gray_hash.update(json.dumps(list(array.shape)).encode())
            gray_hash.update(np.asarray(array, dtype="<i8").tobytes())
            result["grayscale_sha256"] = gray_hash.hexdigest()
            if role in {"tumor", "other"}:
                values = np.unique(array).tolist()
                result["unique_value_count"] = len(values)
                result["unique_values"] = [int(value) for value in values[:20]]
                if not (set(values).issubset({0, 1}) or set(values).issubset({0, 255})):
                    result["issues"].append("nonbinary_mask")
                    return result
                mask = (array > 0).astype("uint8")
                result["foreground_pixels"] = int(mask.sum())
                result["foreground_fraction"] = float(mask.mean())
                if not result["foreground_pixels"]:
                    result["issues"].append("empty_mask")
                resized = np.asarray(Image.fromarray(mask).resize((256, 256), Image.Resampling.NEAREST))
                result["foreground_pixels_256"] = int(resized.sum())
                result["disappears_at_256"] = bool(mask.any() and not resized.any())
                if result["disappears_at_256"]:
                    result["issues"].append("mask_disappears_at_256")
    except (OSError, ValueError, Image.DecompressionBombError):
        result["issues"].append("png_decode_error")
    return result


def read_metadata(path: Path) -> list[dict]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = workbook.worksheets[0].iter_rows(values_only=True)
        header = next(rows)
        if len(set(header)) != len(header) or not set(REQUIRED_COLUMNS).issubset(header):
            raise ValueError("clinical workbook has missing or duplicate required columns")
        selected = {name: header.index(name) for name in REQUIRED_COLUMNS}
        # Deliberately exclude age, diagnosis, symptoms and unnecessary fields.
        return [
            {**{name: row[index] for name, index in selected.items()}, "__source_row__": source_row}
            for source_row, row in enumerate(rows, start=2)
            if any(value is not None for value in row)
        ]
    finally:
        workbook.close()


def _duplicate_groups(inventory: list[dict], key: str, images_only: bool = False) -> list[list[str]]:
    groups: dict[str, list[str]] = defaultdict(list)
    for item in inventory:
        if item.get(key) and (not images_only or item.get("role") == "image"):
            groups[item[key]].append(item["path"])
    return sorted(sorted(paths) for paths in groups.values() if len(paths) > 1)


def _write_csv(path: Path, rows: list[dict], fields: list[str]) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields, extrasaction="ignore")
        writer.writeheader()
        for row in rows:
            writer.writerow({
                key: json.dumps(value, ensure_ascii=False) if isinstance(value, (list, dict)) else value
                for key, value in row.items() if key in fields
            })


def audit_dataset(project_root: Path, raw_dir: Path, output_dir: Path) -> dict:
    """Audit official local files and exclusively create a new output folder."""
    import numpy
    import openpyxl
    import PIL

    project_root, raw_dir, output_dir = (
        Path(path).resolve() for path in (project_root, raw_dir, output_dir)
    )
    # Manifests contain repository-relative paths, never machine-specific paths.
    raw_relative = raw_dir.relative_to(project_root).as_posix()
    output_relative = output_dir.relative_to(project_root).as_posix()
    if output_dir.exists():
        raise FileExistsError("output directory already exists; use a new run identifier")
    if output_dir == raw_dir or raw_dir in output_dir.parents:
        raise ValueError("audit output must be outside the immutable raw directory")
    workbook_path = raw_dir / WORKBOOK_FILENAME
    image_dir = raw_dir / IMAGE_DIRECTORY
    if workbook_path.is_symlink() or image_dir.is_symlink():
        raise ValueError("source workbook and image directory must not be symbolic links")
    if not workbook_path.is_file() or not image_dir.is_dir():
        raise FileNotFoundError("official workbook or image directory is missing")
    metadata = read_metadata(workbook_path)
    if not metadata:
        raise ValueError("clinical workbook contains no case records")
    cases, references = match_rows(metadata)
    inventory = []
    inspected: dict[str, dict] = {}
    global_issues: list[dict] = []
    orphan_paths = []
    for path in sorted(raw_dir.rglob("*")):
        if path.is_symlink():
            global_issues.append({"code": "source_symlink", "path": path.relative_to(project_root).as_posix()})
            continue
        if not path.is_file():
            continue
        relative = path.relative_to(project_root).as_posix()
        item: dict[str, Any] = {
            "path": relative, "bytes": path.stat().st_size,
            "sha256": sha256_file(path), "role": "source_file", "issues": [],
        }
        if path.suffix.lower() == ".png":
            occurrences = references.get(path.name, []) if path.parent == image_dir else []
            roles = sorted({occurrence["role"] for occurrence in occurrences})
            role = roles[0] if len(roles) == 1 else "unreferenced" if not roles else "conflicting"
            item["case_ids"] = [cases[entry["row_index"]]["case_id"] for entry in occurrences]
            if not occurrences:
                orphan_paths.append(relative)
                item["issues"].append("orphan_png")
            details = inspect_png(path, role)
            prior_issues = item["issues"]
            item.update(details)
            item["issues"] = prior_issues + details["issues"]
            if path.parent == image_dir:
                inspected[path.name] = item
        inventory.append(item)
    for name, occurrences in references.items():
        if name not in inspected:
            for occurrence in occurrences:
                cases[occurrence["row_index"]]["issues"].append(f"missing_{occurrence['role']}_file")
    # Duplicate source images require review before any patient split, even
    # though this cannot detect every near duplicate or repeated participant.
    duplicate_file_images = _duplicate_groups(inventory, "sha256", images_only=True)
    duplicate_decoded_images = _duplicate_groups(inventory, "decoded_sha256", images_only=True)
    duplicate_pixel_images = _duplicate_groups(inventory, "grayscale_sha256", images_only=True)
    duplicate_images = {path for group in duplicate_file_images + duplicate_decoded_images + duplicate_pixel_images for path in group}
    for case in cases:
        for role in ("image", "tumor", "other"):
            names = case[f"{role}_filenames"]
            case[f"{role}_paths"] = [(image_dir / name).relative_to(project_root).as_posix() for name in names]
            for name in names:
                details = inspected.get(name)
                if details:
                    case["issues"].extend(f"{role}:{issue}" for issue in details["issues"])
        images = [inspected[name] for name in case["image_filenames"] if name in inspected]
        if len(images) == 1 and "width" in images[0]:
            image = images[0]
            case["width"], case["height"] = image["width"], image["height"]
            case["image_mode"] = image["mode"]
            case["image_grayscale_sha256"] = image.get("grayscale_sha256", "")
            if image["path"] in duplicate_images:
                case["review_flags"].append("duplicate_image_content")
            for role in ("tumor", "other"):
                for name in case[f"{role}_filenames"]:
                    mask = inspected.get(name, {})
                    if "width" in mask and (mask["width"], mask["height"]) != (image["width"], image["height"]):
                        case["issues"].append(f"{role}:image_mask_dimension_mismatch")
            if len(case["tumor_filenames"]) == 1:
                target = inspected.get(case["tumor_filenames"][0], {})
                if "foreground_pixels" in target:
                    case["target_area_px"] = target["foreground_pixels"]
                    case["target_area_fraction"] = target["foreground_fraction"]
                    case["target_area_px_256"] = target["foreground_pixels_256"]
        case["issues"] = sorted(set(case["issues"]))
        case["review_flags"] = sorted(set(case["review_flags"]))
        case["automatic_status"] = "blocked" if case["issues"] else "review_required" if case["review_flags"] else "passed"
    png_items = [item for item in inventory if item["path"].lower().endswith(".png")]
    mask_items = [item for item in png_items if item["role"] in {"tumor", "other"}]
    issue_counts = Counter(issue for case in cases for issue in case["issues"])
    summary = {
        "audit_schema_version": "1.1",
        "raw_directory": raw_relative, "output_directory": output_relative,
        "source_workbook": workbook_path.relative_to(project_root).as_posix(),
        "source_workbook_sha256": sha256_file(workbook_path),
        "case_id_column": "CaseID", "patient_id_basis": PATIENT_ID_BASIS,
        "metadata_columns_retained": list(REQUIRED_COLUMNS),
        "case_count": len(cases),
        "classification_counts": dict(sorted(Counter(case["classification"] for case in cases).items())),
        "automatic_status_counts": dict(sorted(Counter(case["automatic_status"] for case in cases).items())),
        "case_issue_counts": dict(sorted(issue_counts.items())),
        "review_flag_counts": dict(sorted(Counter(flag for case in cases for flag in case["review_flags"]).items())),
        "file_count": len(inventory), "png_count": len(png_items),
        "png_role_counts": dict(sorted(Counter(item["role"] for item in png_items).items())),
        "png_mode_counts": dict(sorted(Counter(item.get("mode", "undecodable") for item in png_items).items())),
        "alpha_handling_counts": dict(sorted(Counter(item.get("alpha_handling", "undecodable") for item in png_items).items())),
        "mask_value_sets": sorted({json.dumps(item["unique_values"]) for item in png_items if "unique_values" in item}),
        "orphan_png_paths": orphan_paths, "global_issues": global_issues,
        "duplicate_png_file_groups": _duplicate_groups(png_items, "sha256"),
        "duplicate_decoded_png_groups": _duplicate_groups(png_items, "decoded_sha256"),
        "duplicate_image_grayscale_groups": duplicate_pixel_images,
        "duplicate_image_decoded_groups": duplicate_decoded_images,
        "duplicate_image_file_groups": duplicate_file_images,
        "masks_disappearing_at_256": [item["path"] for item in png_items if item.get("disappears_at_256")],
        "mask_check_coverage": {
            "masks_total": len(mask_items),
            "masks_with_validated_binary_values": sum("foreground_pixels" in item for item in mask_items),
            "masks_resize_checked": sum("foreground_pixels_256" in item for item in mask_items),
            "masks_not_resize_checked": sum("foreground_pixels_256" not in item for item in mask_items),
            "by_role": {role: {
                "total": sum(item["role"] == role for item in mask_items),
                "validated_binary": sum(item["role"] == role and "foreground_pixels" in item for item in mask_items),
                "resize_checked": sum(item["role"] == role and "foreground_pixels_256" in item for item in mask_items),
            } for role in ("tumor", "other")},
            "interpretation": "An empty disappearance list is not a pass when resize coverage is incomplete. Nonopaque/color encodings remain unresolved; observations below do not authorize a target decoding rule.",
        },
        "encoding_observations": {
            role: {
                "files": sum(item["role"] == role for item in png_items),
                "rgb_channels_equal_files": sum(item["role"] == role and item.get("rgb_channels_equal", False) for item in png_items),
                "rgb_values_only_0_255_files": sum(item["role"] == role and item.get("rgb_values_only_0_255", False) for item in png_items),
                "has_transparency_files": sum(item["role"] == role and item.get("alpha_zero_pixels", 0) > 0 for item in png_items),
                "transparent_color_not_all_zero_paths": [item["path"] for item in png_items if item["role"] == role and item.get("transparent_color_all_zero") is False],
                "partial_alpha_files": [{"path": item["path"], "pixels": item["alpha_partial_pixels"]} for item in png_items if item["role"] == role and item.get("alpha_partial_pixels", 0)],
                "alpha_equals_first_channel_files": sum(item["role"] == role and item.get("alpha_equals_first_channel", False) for item in png_items),
                "opaque_color_all_255_files": sum(item["role"] == role and item.get("opaque_color_all_255", False) for item in png_items),
            } for role in ("image", "tumor", "other")
        },
        "multiple_tumor_case_ids": [case["case_id"] for case in cases if len(case["tumor_filenames"]) > 1],
        "visual_review_status": "pending_visual_review", "training_ready": False,
        "automatic_checks_passed": not issue_counts and not global_issues and not orphan_paths,
        "split_created": False, "training_performed": False,
        "scope_notes": [
            "Automatic consistency checks do not establish that contours are medically correct.",
            "All cases require visual review; passed does not mean ready for training.",
            "Tumor and other masks are kept separate; multiple tumor masks are never merged.",
            "Normal status comes from Classification, never from absent tumor files alone.",
            "256x256 nearest-neighbor checks do not save transformed images or create a data split.",
            "Mask nonemptiness, binary-value validation and resize claims apply only to the explicitly counted coverage; unresolved encodings are not failed clinical cases or automatic exclusions.",
            "Only exact file and decoded-pixel duplicates are tested; near-duplicate review remains outstanding.",
            "No independent training, GPU, TensorFlow or model validation is performed.",
        ],
        "software": {"python": platform.python_version(), "numpy": numpy.__version__, "Pillow": PIL.__version__, "openpyxl": openpyxl.__version__},
    }
    output_dir.mkdir(parents=True, exist_ok=False)
    case_fields = [
        "source_row", "case_id", "patient_id", "image_id", "classification", "margin",
        "image_paths", "tumor_paths", "other_paths", "width", "height", "image_mode",
        "image_grayscale_sha256", "target_area_px", "target_area_fraction", "target_area_px_256",
        "candidate_cohort", "automatic_status", "issues", "review_flags", "visual_review_status", "training_ready",
    ]
    inventory_fields = [
        "path", "bytes", "sha256", "role", "case_ids", "width", "height", "mode", "dtype",
        "alpha_handling", "alpha_min", "alpha_max", "grayscale_min", "grayscale_max",
        "rgb_channels_equal", "rgb_values_only_0_255", "alpha_zero_pixels", "alpha_partial_pixels",
        "transparent_color_all_zero", "alpha_equals_first_channel", "opaque_color_all_255",
        "decoded_sha256", "grayscale_sha256", "unique_value_count", "unique_values",
        "foreground_pixels", "foreground_fraction", "foreground_pixels_256", "disappears_at_256", "issues",
    ]
    _write_csv(output_dir / "cases.csv", cases, case_fields)
    _write_csv(output_dir / "file_inventory.csv", inventory, inventory_fields)
    summary["output_sha256"] = {name: sha256_file(output_dir / name) for name in ("cases.csv", "file_inventory.csv")}
    (output_dir / "audit_summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return summary
